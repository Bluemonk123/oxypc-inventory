"""Shared helpers for the /workid-status tests (2026-10-27).

The page no longer renders its table as <tr id="wo-..."> HTML: rows ship as
JSON in <script id="workidRows"> (routers/workid_status.py _rows_json, short
keys: k=row key, w=WorkID, b=tag, s=stage label, e=engineer, c=completed
date, cd=completed day, d=aging days, st=status, x=scrap, p=PNA, pt=PNA tip).
The export is an .xlsx workbook (Tag Summary / Stage Summary / Staff Summary)
instead of a CSV. These helpers read both, so the tests check the same
behaviour against the current output.
"""
import io
import json
import re

_ROWS_RE = re.compile(r'<script type="application/json" id="workidRows">(.*?)</script>', re.S)


def page_rows(html: str) -> list:
    """Every table row of a rendered /workid-status page, as dicts."""
    m = _ROWS_RE.search(html)
    if not m:
        raise AssertionError("workidRows JSON block not found on the page")
    return json.loads(m.group(1).replace("<\\/", "</"))


def rows_by_workid(html: str) -> dict:
    return {r["w"]: r for r in page_rows(html)}


def export_sheet(content: bytes, sheet: str = "Tag Summary") -> list:
    """One sheet of the /workid-status/export workbook as a list of rows."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True)
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
