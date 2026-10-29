"""Parts Consumption -> Repair Summary export (2026-10-29): one row per tag,
each changed part as its own Part N Name / Part N Cost pair, Total Parts Cost,
Labour / Packing / Freight (Additional Rates, services/repair_cost.py) and
Grand Total. Follows the filter bar and, when given, the table's search."""
import io
import pathlib
import subprocess
import sys
import uuid

from tests.test_iqc_new_user import _login, make_user  # noqa: F401  (fixture)

ROOT = str(pathlib.Path(__file__).resolve().parent.parent)


def _run(src):
    r = subprocess.run([sys.executable, "-c", src], capture_output=True, text=True, cwd=ROOT, timeout=120)
    assert r.returncode == 0, r.stderr[-2000:]
    return r.stdout.strip()


def test_repair_summary_export(app_client, make_user):  # noqa: F811
    from openpyxl import load_workbook
    sfx = uuid.uuid4().hex[:6].upper()
    t1, t2 = f"ITRS{sfx}A", f"ITRS{sfx}B"
    _run(f"""
import asyncio, sys, uuid
sys.path.insert(0, r"{ROOT}")
from sqlalchemy import text
from database import AsyncSessionLocal
from models.device import Device, DeviceStage, StageMovement
from models.spare_parts import SparePart
from models.part_request import PartRequest
from utils.timezone import app_now
async def main():
    async with AsyncSessionLocal() as db:
        lot = (await db.execute(text("select id from lots limit 1"))).scalar()
        parts = []
        for nm, price in (("ITRS-Keyboard-{sfx}", 500), ("ITRS-Battery-{sfx}", 1200), ("ITRS-Hinge-{sfx}", 150)):
            sp = SparePart(id=uuid.uuid4(), part_code=("P" + uuid.uuid4().hex[:10]).upper(), name=nm,
                           category="ITRS", unit_price=price)
            db.add(sp); parts.append(sp)
        await db.flush()
        devs = {{}}
        for bc in ("{t1}", "{t2}"):
            d = Device(id=uuid.uuid4(), barcode=bc, lot_id=lot, current_stage=DeviceStage.ready_to_sale,
                       device_type="Laptop", brand="Dell", model="Latitude 5490", grn_number="ITRSGRN{sfx}")
            db.add(d); devs[bc] = d
        await db.flush()
        db.add(StageMovement(device_id=devs["{t1}"].id, from_stage=DeviceStage.stock_in, to_stage=DeviceStage.l1, moved_by="t"))
        db.add(StageMovement(device_id=devs["{t1}"].id, from_stage=DeviceStage.final_qc,
                             to_stage=DeviceStage.ready_to_sale, moved_by="t"))
        for bc, idx, qty in (("{t1}", 0, 1), ("{t1}", 1, 1), ("{t1}", 2, 2), ("{t2}", 0, 1)):
            sp = parts[idx]
            db.add(PartRequest(device_id=devs[bc].id, part_id=sp.id, part_name=sp.name, status="received",
                               qty_requested=qty, qty_handed_over=qty, engineer_name="IT Eng {sfx}",
                               actioned_at=app_now()))
        await db.commit()
asyncio.run(main())
""")
    try:
        username, password = make_user("admin")
        _login(app_client, username, password)
        html = app_client.get("/spare-parts").text
        assert 'id="pcExportRepairSummaryBtn"' in html
        csrf = app_client.cookies.get("csrf_token") or "dummy"

        def export(**extra):
            r = app_client.post("/spare-parts/repair-summary", data={"csrf_token": csrf, "category": "ITRS", **extra},
                                headers={"X-CSRF-Token": csrf})
            assert r.status_code == 200, r.text[:300]
            ws = load_workbook(io.BytesIO(r.content)).active
            head = [c.value for c in ws[1]]
            rows = {r[0]: dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
            return head, rows

        head, rows = export()
        # busiest tag has 3 parts -> 3 Name / Cost pairs
        assert head[:3] == ["Tag Number", "GRN", "Invoice Number"]
        assert "Part 3 Name" in head and "Part 3 Cost" in head and "Part 4 Name" not in head
        assert head[-5:] == ["Total Parts Cost", "Labour Cost", "Packing Cost", "Freight Cost", "Grand Total"]
        a = rows[t1]
        assert a["GRN"] == f"ITRSGRN{sfx}" and a["Device Type"] == "Laptop" and a["Engineer Name"] == f"IT Eng {sfx}"
        assert a["Final QC Pass Date"] is not None
        costs = sorted([a["Part 1 Cost"], a["Part 2 Cost"], a["Part 3 Cost"]])
        assert costs == [300, 500, 1200]                       # hinge x2 = 300
        assert a["Total Parts Cost"] == 2000
        assert a["Grand Total"] == round(2000 + a["Labour Cost"] + a["Packing Cost"] + a["Freight Cost"], 2)
        assert a["Labour Cost"] > 0                             # worked on (L1) laptop
        b = rows[t2]
        assert b["Total Parts Cost"] == 500 and b["Part 2 Name"] in (None, "") and b["Labour Cost"] == 0
        # the table's search narrowed to one tag
        _h, rows = export(tags_given="1", tags=t2)
        assert set(rows) == {t2}

        # Consumption Summary (2026-10-29): Tags Summary + Parts Summary in one file
        assert "Consumption Summary" in html and "/spare-parts/consumption-summary" in html
        r = app_client.post("/spare-parts/consumption-summary", data={"csrf_token": csrf, "category": "ITRS"},
                            headers={"X-CSRF-Token": csrf})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Tag Summary", "Parts Summary"]
        tag_rows = [x for x in wb["Tag Summary"].iter_rows(min_row=2, values_only=True)]
        assert len(tag_rows) == 4 and {x[0] for x in tag_rows} == {t1, t2}
        ps = {x[0]: x for x in wb["Parts Summary"].iter_rows(min_row=2, values_only=True)}
        kb = ps[f"ITRS-Keyboard-{sfx}"]
        assert kb[1] == 2 and kb[2] == 2 and kb[-1] == 1000      # 2 used, on 2 tags, 2 x 500

        # WorkID Status no longer has a Repair Summary (removed 2026-10-29)
        assert "/workid-status/repair-summary" not in app_client.get("/workid-status").text
    finally:
        _run(f"""
import asyncio, sys
sys.path.insert(0, r"{ROOT}")
from sqlalchemy import text
from database import AsyncSessionLocal
async def main():
    async with AsyncSessionLocal() as db:
        ids = "(select id from devices where barcode like 'ITRS{sfx}%')"
        await db.execute(text(f"delete from part_requests where device_id in {{ids}}"))
        await db.execute(text(f"delete from stage_movements where device_id in {{ids}}"))
        await db.execute(text("delete from devices where barcode like 'ITRS{sfx}%'"))
        await db.execute(text("delete from spare_parts where name like 'ITRS-%{sfx}'"))
        await db.commit()
asyncio.run(main())
""")
